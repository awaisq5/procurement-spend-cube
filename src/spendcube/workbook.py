"""The cube as one Excel workbook, for category managers who work in Excel.

Same numbers as the CSV extracts, with euro, percentage and date formats, a
frozen header row, filters and readable column widths on every sheet.
"""

from __future__ import annotations

from pathlib import Path

import pandas as pd
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.worksheet import Worksheet

from .analytics import Cube
from .pipeline import Result

EUR = '#,##0 "€";-#,##0 "€"'
PCT = "0%"
DATE = "yyyy-mm-dd"
HEADER_FILL = PatternFill("solid", fgColor="1F3B2D")
HEADER_FONT = Font(bold=True, color="FFFFFF")

# Column name -> number format. Anything not listed keeps Excel's default.
FORMATS = {
    "spend": EUR,
    "amount_eur": EUR,
    "at_stake": EUR,
    "current": EUR,
    "prior": EUR,
    "change_eur": EUR,
    "top_share": PCT,
    "group_top_share": PCT,
    "on_contract_share": PCT,
    "change": PCT,
    "confidence": PCT,
    "posting_date": DATE,
}

LINE_COLUMNS = [
    "po_number",
    "posting_date",
    "supplier_id",
    "supplier",
    "supplier_name",
    "supplier_country",
    "description",
    "amount",
    "currency",
    "amount_eur",
    "category_id",
    "category",
    "category_group",
    "confidence",
    "source",
    "on_contract",
]


def _summary(cube: Cube) -> pd.DataFrame:
    t = cube.totals
    rows = [
        ("Period", f"{t['first_date']:%Y-%m-%d} to {t['last_date']:%Y-%m-%d}", None),
        ("Spend", t["spend"], EUR),
        ("Lines", t["lines"], None),
        ("Supplier spellings", t["raw_names"], None),
        ("Suppliers after matching", t["suppliers"], None),
        ("Supplier groups", t["groups"], None),
        ("Spend classified", t["classified_share"], PCT),
        ("Spend on contract", t["on_contract_share"], PCT),
        ("Spend outside contracts where one exists", t["maverick_spend"], EUR),
        ("Suppliers for 80% of spend", t["suppliers_for_80"], None),
        ("Suppliers in the last 5% of spend", t["tail_suppliers"], None),
        ("Categories growing off contract", t.get("yoy_rising_off_contract", 0), None),
    ]
    df = pd.DataFrame([(a, b) for a, b, _ in rows], columns=["measure", "value"])
    df.attrs["row_formats"] = [f for _, _, f in rows]
    return df


def _shown_width(value: object, fmt: str | None) -> int:
    """Roughly how many characters Excel shows for a value in a given format."""
    if isinstance(value, float) and fmt == EUR:
        return len(f"{value:,.0f} €")
    if isinstance(value, float) and fmt == PCT:
        return 5
    if isinstance(value, float):
        return min(len(f"{value:.4g}"), 12)
    return len(str(value))


def _style(ws: Worksheet, df: pd.DataFrame) -> None:
    for cell in ws[1]:
        cell.fill, cell.font = HEADER_FILL, HEADER_FONT
        cell.alignment = Alignment(vertical="center")
    ws.freeze_panes = "A2"
    if len(df):
        ws.auto_filter.ref = ws.dimensions
    for i, col in enumerate(df.columns, start=1):
        letter = get_column_letter(i)
        fmt = FORMATS.get(col)
        if fmt:
            for (cell,) in ws.iter_rows(min_row=2, min_col=i, max_col=i):
                cell.number_format = fmt
        sample = [len(str(col))] + [_shown_width(v, fmt) for v in df[col].head(200)]
        ws.column_dimensions[letter].width = min(max(sample) + 2, 60)
    for row, fmt in enumerate(df.attrs.get("row_formats", []), start=2):
        if fmt:
            ws.cell(row=row, column=2).number_format = fmt


def write_workbook(result: Result, cube: Cube, path: Path) -> Path:
    groups = cube.groups.assign(member_names=cube.groups.member_names.map("; ".join))
    review = result.review.assign(posting_date=pd.to_datetime(result.review.posting_date).dt.date)
    lines = result.lines[LINE_COLUMNS].assign(posting_date=pd.to_datetime(result.lines.posting_date).dt.date)
    sheets = {
        "Summary": _summary(cube),
        "Categories": cube.categories,
        "Year on year": cube.year_on_year,
        "Suppliers": cube.suppliers,
        "Groups": groups,
        "Review queue": review,
        "Classified lines": lines,
    }
    with pd.ExcelWriter(path, engine="openpyxl") as xl:
        for name, df in sheets.items():
            if name == "Year on year" and df.empty:
                continue
            df.to_excel(xl, sheet_name=name, index=False)
            _style(xl.sheets[name], df)
    return path
